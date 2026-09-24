"""Generate INSERT-ALL-43-VERBATIM-FROM-DOCX.sql from the verbatim DOCX table."""
from __future__ import annotations

import re
from pathlib import Path

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

ROOT = Path(__file__).resolve().parent
DOCX = Path.home() / "Downloads" / "OSG_DOST_Cases_1_to_43_VERBATIM_Status_Remarks (2).docx"
OUTPUT = ROOT / "INSERT-ALL-43-VERBATIM-FROM-DOCX.sql"

# Case metadata aligned with OSG report numbering 1–43
CASE_META = {
    1: {"code": "c-usigan-001", "case_type": "Civil Case", "status": "Ongoing", "filing_date": "2023-01-01", "last_updated": "2025-12-05"},
    2: {"code": "c-baculi-001", "case_type": "Civil Case", "status": "Pending", "filing_date": "2022-01-01", "last_updated": "2024-01-01"},
    3: {"code": "c-lasam-lakas-001", "case_type": "Civil Case", "status": "Ongoing", "filing_date": "2022-01-01", "last_updated": "2024-12-03"},
    4: {"code": "c-collado-001", "case_type": "Civil Case", "status": "Archived", "filing_date": "2021-01-01", "last_updated": "2021-09-28"},
    5: {"code": "c-yapit-001", "case_type": "Civil Case", "status": "Ongoing", "filing_date": "2023-01-01", "last_updated": "2024-11-26"},
    6: {"code": "c-carabacan-001", "case_type": "Civil Case", "status": "Ongoing", "filing_date": "2023-01-01", "last_updated": "2024-09-16"},
    7: {"code": "c-magana-001", "case_type": "Civil Case", "status": "Ongoing", "filing_date": "2023-01-01", "last_updated": "2025-05-21"},
    8: {"code": "c-mangaoang-001", "case_type": "Civil Case", "status": "Ongoing", "filing_date": "2023-01-01", "last_updated": "2025-11-06"},
    9: {"code": "c-chavez-001", "case_type": "Civil Case", "status": "Ongoing", "filing_date": "2023-01-01", "last_updated": "2024-07-31"},
    10: {"code": "c-ranjo-001", "case_type": "Special Proceeding", "status": "Ongoing", "filing_date": "2024-01-01", "last_updated": "2025-11-13"},
    11: {"code": "c-coloma-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    12: {"code": "c-rosendo-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2022-05-05", "last_updated": "2022-05-05"},
    13: {"code": "c-mamauag-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    14: {"code": "c-colobong-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    15: {"code": "c-pagunuran-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    16: {"code": "c-massalang-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    17: {"code": "c-cruz-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    18: {"code": "c-tumbali-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    19: {"code": "c-abalos-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2022-11-08", "last_updated": "2022-11-08"},
    20: {"code": "c-cubacub-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    21: {"code": "c-villanueva-001", "case_type": "Small Claims", "status": "Pending", "filing_date": "2022-11-08", "last_updated": "2022-11-08"},
    22: {"code": "c-galapon-001", "case_type": "Small Claims", "status": "Pending", "filing_date": "2023-03-15", "last_updated": "2023-03-15"},
    23: {"code": "c-tomas-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    24: {"code": "c-domingo-001", "case_type": "Small Claims", "status": "Pending", "filing_date": "2023-07-31", "last_updated": "2023-08-07"},
    25: {"code": "c-uy-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    26: {"code": "c-duerme-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    27: {"code": "c-khong-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-02-22", "last_updated": "2023-02-22"},
    28: {"code": "c-garcia-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    29: {"code": "c-dimaculangan-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    30: {"code": "c-mallare-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-03-30", "last_updated": "2023-03-31"},
    31: {"code": "c-lamiere-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-11-10", "last_updated": "2023-11-10"},
    32: {"code": "c-pua-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-11-10", "last_updated": "2023-11-10"},
    33: {"code": "c-delossantos-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-11-10", "last_updated": "2023-11-10"},
    34: {"code": "c-barien-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-11-10", "last_updated": "2023-11-10"},
    35: {"code": "c-navis-001", "case_type": "Small Claims", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    36: {"code": "c-pinzon-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-05-15", "last_updated": "2023-05-15"},
    37: {"code": "c-pablo-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-01-01", "last_updated": "2023-01-01"},
    38: {"code": "c-base-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-05-23", "last_updated": "2023-05-23"},
    39: {"code": "c-montilla-001", "case_type": "Small Claims", "status": "Pending", "filing_date": "2023-08-18", "last_updated": "2023-08-18"},
    40: {"code": "c-paraiso-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2023-10-25", "last_updated": "2023-10-25"},
    41: {"code": "c-lasam-kikos-001", "case_type": "Small Claims", "status": "Pending", "filing_date": "2023-08-18", "last_updated": "2023-08-18"},
    42: {"code": "c-guzman-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2019-01-01", "last_updated": "2019-01-01"},
    43: {"code": "c-andres-001", "case_type": "Pre-litigation", "status": "Pending", "filing_date": "2024-01-29", "last_updated": "2024-01-29"},
}

PERIOD_RE = re.compile(
    r"^(JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER|APRIL|JUNE)\b",
    re.I,
)
EMPTY_REMARK_PLACEHOLDER = (
    "No status/remarks text is recorded in the supplied Status/Remarks cells for this entry."
)
BULLET_PREFIX = re.compile(r"^[\u2022\uf0b7\u00b7\u2023\u2043\u2219]\s*")
UNASSIGNED = {"—", "-", "–", "", "—"}


def normalize_text(value: str) -> str:
    text = value.replace("\r", "").replace("\u00a0", " ").strip()
    text = text.replace("\ufffd", "–")
    text = re.sub(r"\s+", " ", text)
    return text


def normalize_remark(line: str) -> str | None:
    text = line.strip()
    if not text:
        return None
    text = BULLET_PREFIX.sub("", text)
    if text.startswith("o "):
        text = f"  {text[2:].strip()}"
    return text


def finalize_remarks(remarks: list[str]) -> list[str]:
    if remarks:
        return remarks
    return [EMPTY_REMARK_PLACEHOLDER]


def build_parties(title: str, num: int) -> str:
    if num <= 10 and title.startswith("Republic of the Philippines"):
        if " v. " in title:
            defendant = title.split(" v. ", 1)[1]
            return f"Plaintiff: Republic of the Philippines, rep. by DOST-RO2; Defendant: {defendant}"
        return f"Petitioners: {title.split(' of: ', 1)[-1]}"
    vs_name = title.split("(")[0].replace("Mr. ", "").replace("Mrs. ", "").replace("Ms. ", "").strip()
    return f"Republic of the Philippines, rep. by DOST-RO2 vs. {vs_name}"


def sql_str(value: str | None) -> str:
    if value is None or value in UNASSIGNED:
        return "NULL"
    escaped = value.replace("'", "''")
    return f"'{escaped}'"


def merge_period_headers(raw_remarks: list[str]) -> list[str]:
    remarks: list[str] = []
    pending_period: str | None = None
    for remark in raw_remarks:
        if PERIOD_RE.match(remark) and "–" in remark:
            pending_period = remark
            continue
        if pending_period:
            remarks.append(f"{pending_period}: {remark}")
            pending_period = None
        else:
            remarks.append(remark)
    if pending_period:
        remarks.append(pending_period)
    return remarks


def parse_cases_from_docx(docx_path: Path) -> list[dict]:
    if Document is None:
        raise SystemExit("Install python-docx: pip install python-docx")

    doc = Document(docx_path)
    if not doc.tables:
        raise SystemExit(f"No tables found in {docx_path}")

    table = doc.tables[0]
    cases: list[dict] = []

    for row in table.rows[1:]:
        cells = [normalize_text(cell.text) for cell in row.cells]
        if not cells or not cells[0].isdigit():
            continue

        num = int(cells[0])
        title = cells[1]
        case_number = cells[2] if cells[2] not in UNASSIGNED else None
        court = cells[3] if cells[3] not in UNASSIGNED else None

        raw_remarks: list[str] = []
        for line in row.cells[4].text.replace("\r", "").split("\n"):
            remark = normalize_remark(line)
            if remark:
                raw_remarks.append(remark)

        cases.append(
            {
                "num": num,
                "title": title,
                "case_number": case_number,
                "court": court,
                "remarks": finalize_remarks(merge_period_headers(raw_remarks)),
            }
        )

    return cases


def parse_cases(lines: list[str]) -> list[dict]:
    start = next(i for i, line in enumerate(lines) if line == "1")
    cases: list[dict] = []
    i = start
    while i < len(lines):
        if not re.fullmatch(r"\d{1,2}", lines[i]):
            i += 1
            continue
        num = int(lines[i])
        title = lines[i + 1]
        case_number = lines[i + 2]
        court = lines[i + 3]
        i += 4
        raw_remarks: list[str] = []
        while i < len(lines) and not re.fullmatch(r"\d{1,2}", lines[i]):
            remark = normalize_remark(lines[i])
            if remark:
                raw_remarks.append(remark)
            i += 1

        # Merge period headers with the next bullet (matches report table style)
        cases.append(
            {
                "num": num,
                "title": title,
                "case_number": None if case_number in UNASSIGNED else case_number,
                "court": None if court in UNASSIGNED else court,
                "remarks": finalize_remarks(merge_period_headers(raw_remarks)),
            }
        )
    return cases


def generate_sql(cases: list[dict]) -> str:
    parts: list[str] = [
        "-- =============================================================================",
        "-- OSG DOST TASK FORCE — DELETE + MERGE ALL 43 CASES + VERBATIM REMARKS",
        "-- Source: OSG_DOST_Cases_1_to_43_VERBATIM_Status_Remarks (2).docx",
        "-- Generated by supabase/generate_verbatim_sql.py",
        "--",
        "-- One file, one run:",
        "--   1) DELETE old/extra cases",
        "--   2) MERGE (upsert) official cases 1–43 by code",
        "--   3) REPLACE all verbatim Status/Remarks",
        "-- Paste into Supabase → SQL Editor → Run",
        "-- =============================================================================",
        "",
        "BEGIN;",
        "",
        "ALTER TABLE public.cases ADD COLUMN IF NOT EXISTS report_case_number integer;",
        "",
    ]

    codes = [CASE_META[n]["code"] for n in range(1, 44)]
    code_list = ", ".join(f"'{c}'" for c in codes)
    parts.extend(
        [
            "-- =============================================================================",
            "-- PHASE 1: DELETE previous / extra cases",
            "-- =============================================================================",
            "",
            "-- Remove child rows for cases that are NOT in the official 43 codes",
            "DELETE FROM public.case_status_updates",
            "WHERE case_id IN (SELECT id FROM public.cases WHERE code NOT IN (",
            f"  {code_list}",
            "));",
            "",
            "DELETE FROM public.case_activity",
            "WHERE case_id IN (SELECT id FROM public.cases WHERE code NOT IN (",
            f"  {code_list}",
            "));",
            "",
            "DELETE FROM public.case_files",
            "WHERE case_id IN (SELECT id FROM public.cases WHERE code NOT IN (",
            f"  {code_list}",
            "));",
            "",
            "-- Delete extra case rows (old test data, duplicates, wrong codes)",
            "DELETE FROM public.cases",
            "WHERE code NOT IN (",
            f"  {code_list}",
            ");",
            "",
            "-- Delete duplicate report numbers if any slipped through (keep official code only)",
            "DELETE FROM public.cases c",
            "USING public.cases d",
            "WHERE c.report_case_number = d.report_case_number",
            "  AND c.report_case_number BETWEEN 1 AND 43",
            "  AND c.id <> d.id",
            f"  AND c.code NOT IN ({code_list});",
            "",
            "-- Clear remarks on the official 43 before fresh verbatim insert",
            "DELETE FROM public.case_status_updates",
            f"WHERE case_id IN (SELECT id FROM public.cases WHERE code IN ({code_list}));",
            "",
            "-- =============================================================================",
            "-- PHASE 2: MERGE (UPSERT) cases 1–43",
            "-- ON CONFLICT (code) DO UPDATE = merge into existing row if code already exists",
            "-- =============================================================================",
            "",
        ]
    )

    total_remarks = 0
    for case in cases:
        num = case["num"]
        meta = CASE_META[num]
        code = meta["code"]
        parties = build_parties(case["title"], num)
        remarks = case["remarks"]
        total_remarks += len(remarks)
        first_remark = remarks[0] if remarks else None

        parts.extend(
            [
                f"-- CASE #{num}: {case['title'][:70]}",
                "INSERT INTO public.cases (",
                "  code, case_title, case_number, case_type, court, parties, status,",
                "  remarks, report_case_number, filing_date, last_updated",
                ") VALUES (",
                f"  '{code}',",
                f"  {sql_str(case['title'])},",
                f"  {sql_str(case['case_number'])},",
                f"  {sql_str(meta['case_type'])},",
                f"  {sql_str(case['court'])},",
                f"  {sql_str(parties)},",
                f"  '{meta['status']}'::public.case_status,",
                f"  {sql_str(first_remark)},",
                f"  {num},",
                f"  '{meta['filing_date']}',",
                f"  '{meta['last_updated']}'",
                ") ON CONFLICT (code) DO UPDATE SET",
                "  case_title = EXCLUDED.case_title,",
                "  case_number = EXCLUDED.case_number,",
                "  case_type = EXCLUDED.case_type,",
                "  court = EXCLUDED.court,",
                "  parties = EXCLUDED.parties,",
                "  status = EXCLUDED.status,",
                "  remarks = EXCLUDED.remarks,",
                "  report_case_number = EXCLUDED.report_case_number,",
                "  filing_date = EXCLUDED.filing_date,",
                "  last_updated = EXCLUDED.last_updated;",
                "",
            ]
        )

        if remarks:
            value_rows = ",\n".join(
                f"  ({idx}, {sql_str(remark)})" for idx, remark in enumerate(remarks, start=1)
            )
            parts.extend(
                [
                    "-- PHASE 3: INSERT verbatim Status/Remarks",
                    "INSERT INTO public.case_status_updates (case_id, sort_order, body)",
                    "SELECT c.id, v.sort_order, v.body",
                    "FROM public.cases c",
                    f"CROSS JOIN (VALUES\n{value_rows}\n) AS v(sort_order, body)",
                    f"WHERE c.code = '{code}';",
                    "",
                ]
            )

    parts.extend(
        [
            "-- =============================================================================",
            "-- VERIFICATION",
            "-- =============================================================================",
            "SELECT report_case_number, code, LEFT(case_title, 60) AS case_title, status,",
            "  (SELECT COUNT(*) FROM public.case_status_updates u WHERE u.case_id = c.id) AS remark_count",
            "FROM public.cases c",
            "WHERE report_case_number BETWEEN 1 AND 43",
            "ORDER BY report_case_number;",
            "",
            "SELECT COUNT(*) AS total_cases FROM public.cases WHERE report_case_number BETWEEN 1 AND 43;",
            "SELECT COUNT(*) AS total_remarks FROM public.case_status_updates u",
            "JOIN public.cases c ON c.id = u.case_id",
            "WHERE c.report_case_number BETWEEN 1 AND 43;",
            "",
            f"-- Expected: 43 cases, {total_remarks} total remark rows",
            "COMMIT;",
            "",
        ]
    )
    return "\n".join(parts)


def main() -> None:
    if not DOCX.exists():
        raise SystemExit(f"DOCX not found: {DOCX}")

    cases = parse_cases_from_docx(DOCX)
    if len(cases) != 43:
        raise SystemExit(f"Expected 43 cases, parsed {len(cases)}")
    missing = [n for n in range(1, 44) if not any(c["num"] == n for c in cases)]
    if missing:
        raise SystemExit(f"Missing case numbers: {missing}")

    sql = generate_sql(cases)
    OUTPUT.write_text(sql, encoding="utf-8")
    remark_total = sum(len(c["remarks"]) for c in cases)
    print(f"Source: {DOCX.name}")
    print(f"Wrote {OUTPUT.name}: 43 cases, {remark_total} remarks")
    for case in cases:
        print(f"  Case {case['num']:2d}: {len(case['remarks']):3d} remarks")


if __name__ == "__main__":
    main()
